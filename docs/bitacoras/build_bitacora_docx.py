from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "bitacora-2026-08-24-a-2026-09-22.md"
OUTPUT = ROOT / "bitacora-itam-2026-08-24-a-2026-09-22.docx"

NAVY = "123B5D"
TEAL = "0C7182"
INK = "17212B"
MUTED = "52616F"
GRID = "D9E1E8"
PALE = "F3F7FA"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), GRID)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_run(run, size=10.1, color=INK, bold=False, italic=False, font="Arial"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_inline(paragraph, text, size=10.1, color=INK):
    token = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))")
    position = 0
    for match in token.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            style_run(run, size=size, color=color)
        value = match.group(0)
        if value.startswith("**"):
            run = paragraph.add_run(value[2:-2])
            style_run(run, size=size, color=color, bold=True)
        elif value.startswith("`"):
            run = paragraph.add_run(value[1:-1])
            style_run(run, size=size - 0.2, color=NAVY, font="Consolas")
        else:
            run = paragraph.add_run(value[1:-1])
            style_run(run, size=size, color=color, italic=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        style_run(run, size=size, color=color)


def set_paragraph_defaults(paragraph, after=5, before=0, line=1.08):
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(after)
    fmt.space_before = Pt(before)
    fmt.line_spacing = line
    fmt.widow_control = True


def add_md_table(document, rows):
    clean_rows = []
    for line in rows:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            continue
        clean_rows.append(cells)
    if not clean_rows:
        return
    col_count = max(map(len, clean_rows))
    table = document.add_table(rows=len(clean_rows), cols=col_count)
    table.autofit = False
    widths = [Inches(0.85), Inches(1.6), Inches(4.55)]
    if col_count != 3:
        widths = [Inches(7.1 / col_count)] * col_count
    for row_index, values in enumerate(clean_rows):
        row = table.rows[row_index]
        if row_index == 0:
            set_repeat_table_header(row)
        for col_index in range(col_count):
            cell = row.cells[col_index]
            cell.width = widths[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE)
            paragraph = cell.paragraphs[0]
            set_paragraph_defaults(paragraph, after=0, line=1.0)
            text = values[col_index] if col_index < len(values) else ""
            add_inline(
                paragraph,
                text,
                size=8.7 if row_index else 9.0,
                color="FFFFFF" if row_index == 0 else INK,
            )
            for run in paragraph.runs:
                if row_index == 0:
                    run.bold = True
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = widths[idx]
    set_table_borders(table)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(document, alt, relative_path):
    path = (ROOT / relative_path).resolve()
    if path.name == "hitos-itam-2026.svg":
        path = path.with_suffix(".png")
    if not path.exists():
        raise FileNotFoundError(path)
    if "hitos-itam" in path.name:
        width = Inches(4.75)
    elif "referencia-kpi" in path.name:
        width = Inches(4.35)
    else:
        width = Inches(6.45)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_defaults(paragraph, after=2, before=4, line=1.0)
    shape = paragraph.add_run().add_picture(str(path), width=width)
    shape._inline.docPr.set("descr", alt)
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_defaults(caption, after=8, line=1.0)
    run = caption.add_run(alt)
    style_run(run, size=8.8, color=MUTED, italic=True)


def set_page_field(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instruction, fld_char_separate, text, fld_char_end])
    style_run(run, size=8, color=MUTED)


def configure_document(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.footer_distance = Inches(0.35)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.1)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = document.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")
    title.paragraph_format.space_after = Pt(7)
    title.paragraph_format.keep_with_next = True

    for name, size, before, after in (("Heading 1", 17, 14, 6), ("Heading 2", 13, 11, 5), ("Heading 3", 11.5, 9, 4)):
        style = document.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_defaults(footer, after=0, line=1.0)
    label = footer.add_run("Aguas San Isidro  |  ITAM  |  ")
    style_run(label, size=8, color=MUTED)
    set_page_field(footer)

    props = document.core_properties
    props.title = "Bitácora detallada de trabajo ITAM"
    props.subject = "Registro diario de actividades y hitos, 24 de agosto a 22 de septiembre de 2026"
    props.author = ""
    props.last_modified_by = ""


def main():
    replace_existing = "--replace" in sys.argv[1:]
    if OUTPUT.exists() and not replace_existing:
        raise FileExistsError(f"Refusing to overwrite existing deliverable: {OUTPUT}")
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)

    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            block = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index].strip())
                index += 1
            add_md_table(document, block)
            continue
        image_match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image_match:
            add_image(document, image_match.group(1), image_match.group(2))
            index += 1
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            add_inline(paragraph, line[2:].strip(), size=25, color="000000")
            set_paragraph_defaults(paragraph, after=7, line=1.0)
            index += 1
            continue
        if line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_inline(paragraph, line[4:].strip(), size=13, color="000000")
            index += 1
            continue
        if line.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1")
            add_inline(paragraph, line[3:].strip(), size=17, color="000000")
            index += 1
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            set_paragraph_defaults(paragraph, after=3.5, line=1.04)
            add_inline(paragraph, line[2:].strip(), size=9.9)
            index += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            set_paragraph_defaults(paragraph, after=3.5, line=1.04)
            add_inline(paragraph, numbered.group(2), size=9.9)
            index += 1
            continue

        paragraph = document.add_paragraph()
        set_paragraph_defaults(paragraph, after=5, line=1.08)
        if line.startswith("*") and line.endswith("*") and not line.startswith("**"):
            run = paragraph.add_run(line.strip("*"))
            style_run(run, size=8.8, color=MUTED, italic=True)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            add_inline(paragraph, line, size=10.1)
            if line.startswith("**Período:"):
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor.from_string(MUTED)
        index += 1

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
